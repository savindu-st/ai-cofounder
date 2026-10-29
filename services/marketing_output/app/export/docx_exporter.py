"""Editable Word exporter built with python-docx."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Iterable, List

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from shared.contracts.roadmap import StartupRoadmap


def _money(value: Any) -> str:
    try:
        return f"${float(value):,.0f}"
    except (TypeError, ValueError):
        return "$0"


def _clean(value: Any) -> str:
    return "" if value is None else str(value)


def _add_bullets(document: Document, items: Iterable[Any]) -> None:
    values = [str(item) for item in (items or []) if str(item).strip()]
    if not values:
        document.add_paragraph("Not provided")
        return
    for item in values:
        document.add_paragraph(item, style="List Bullet")


def _add_table(document: Document, headers: List[str], rows: List[List[str]], widths: List[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].text = ""
        run = header_cells[index].paragraphs[0].add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = _clean(value)
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    if widths:
        for index, width in enumerate(widths):
            for row in table.rows:
                row.cells[index].width = width
    return table


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(30, 58, 138)


def generate_docx(roadmap: StartupRoadmap) -> bytes:
    """Return an editable Word document as bytes."""
    if not isinstance(roadmap, StartupRoadmap):
        raise TypeError("generate_docx expects a StartupRoadmap instance")

    founder = roadmap.founder_input
    revenue = roadmap.revenue_estimation
    market = roadmap.market_research
    marketing = roadmap.marketing_plan
    business_canvas = roadmap.business_model.canvas

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(42)
    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10)

    document.core_properties.title = f"{roadmap.venture_name} - Executive Roadmap"
    document.core_properties.author = "AI Co-Founder"
    document.core_properties.subject = "Executive startup roadmap and business plan"

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("EXECUTIVE STARTUP ROADMAP\n")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(71, 85, 105)
    run = title.add_run(roadmap.venture_name)
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(23, 37, 84)

    subtitle = document.add_paragraph(roadmap.idea_analysis.refined_value_proposition)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True

    metadata = document.add_table(rows=0, cols=2)
    metadata.style = "Table Grid"
    for label, value in [
        ("Founder idea", founder.startup_idea),
        ("Target market", founder.target_market),
        ("Geography", founder.target_geography),
        ("Budget", _money(founder.budget)),
        ("Timeline", f"{founder.timeline_months} months"),
        ("Overall confidence", f"{roadmap.overall_confidence_score:.0%}"),
        ("Generated", datetime.now(timezone.utc).strftime("%B %d, %Y %H:%M UTC")),
    ]:
        cells = metadata.add_row().cells
        cells[0].text = label
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = _clean(value)

    document.add_page_break()
    _add_heading(document, "Executive Summary", 1)
    document.add_paragraph(roadmap.executive_summary)

    _add_table(
        document,
        ["Signal", "Snapshot"],
        [
            ["Confidence", f"{roadmap.overall_confidence_score:.0%}"],
            ["TAM / SAM / SOM", f"{_money(market.tam_estimate)} / {_money(market.sam_estimate)} / {_money(market.som_estimate)}"],
            ["Pricing strategy", revenue.pricing_strategy],
            ["Priority channels", ", ".join(marketing.priority_channels) or "Founder-led outreach"],
            ["Archetype", marketing.venture_archetype or "Not classified"],
            ["Budget allocated", _money(sum((marketing.budget_allocation or {}).values()))],
        ],
    )

    if revenue.unit_economics:
        unit = revenue.unit_economics
        _add_heading(document, "Unit Economics", 2)
        _add_table(
            document,
            ["ARPU", "Gross margin", "CAC", "LTV", "LTV:CAC", "CAC payback"],
            [[
                _money(unit.arpu),
                f"{unit.gross_margin:.1%}",
                _money(unit.cac),
                _money(unit.ltv),
                f"{unit.ltv_cac_ratio:.1f}x",
                f"{unit.cac_payback_months:.1f} mo",
            ]],
        )

    _add_heading(document, "Business Model Canvas", 1)
    bmc_blocks = [
        ("Key Partners", business_canvas.key_partners),
        ("Key Activities", business_canvas.key_activities),
        ("Value Propositions", business_canvas.value_propositions),
        ("Key Resources", business_canvas.key_resources),
        ("Customer Relationships", business_canvas.customer_relationships),
        ("Customer Segments", business_canvas.customer_segments),
        ("Channels", business_canvas.channels),
        ("Cost Structure", business_canvas.cost_structure),
        ("Revenue Streams", business_canvas.revenue_streams),
    ]
    bmc_table = document.add_table(rows=3, cols=3)
    bmc_table.style = "Table Grid"
    for index, (block_title, items) in enumerate(bmc_blocks):
        cell = bmc_table.cell(index // 3, index % 3)
        cell.text = ""
        heading = cell.paragraphs[0]
        run = heading.add_run(block_title)
        run.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
        for item in items or ["Not provided"]:
            cell.add_paragraph(str(item), style="List Bullet")

    _add_heading(document, "Competitor Landscape", 1)
    if market.competitor_landscape:
        _add_table(
            document,
            ["Competitor", "Description", "Strengths", "Weaknesses"],
            [
                [
                    competitor.name,
                    competitor.description,
                    ", ".join(competitor.strengths) or "-",
                    ", ".join(competitor.weaknesses) or "-",
                ]
                for competitor in market.competitor_landscape[:8]
            ],
        )
    else:
        document.add_paragraph("No direct competitors were identified in the research output.")
    if market.source_urls:
        _add_heading(document, "Supporting Sources", 2)
        _add_bullets(document, market.source_urls[:10])

    _add_heading(document, "12-Month Financial Projections", 1)
    _add_table(
        document,
        ["Scenario", "Month 12 revenue", "Month 12 customers", "Break-even", "Ending runway"],
        [
            [
                scenario.scenario_name.title(),
                _money(scenario.month_12_revenue),
                f"{scenario.month_12_customers:,}",
                f"Month {scenario.break_even_month}" if scenario.break_even_month else "Not reached",
                f"{scenario.ending_runway_months:.1f} mo" if scenario.ending_runway_months is not None else "-",
            ]
            for scenario in revenue.scenarios
        ],
    )
    _add_heading(document, "Model Assumptions", 2)
    params = revenue.parameters
    _add_table(
        document,
        ["Monthly price", "Starting customers", "Growth", "Churn", "COGS / unit", "Fixed costs"],
        [[
            _money(params.monthly_price),
            f"{params.starting_customers:,}",
            f"{params.monthly_growth_rate:.1%}",
            f"{params.monthly_churn_rate:.1%}",
            _money(params.cogs_per_unit),
            _money(params.fixed_monthly_costs),
        ]],
    )
    if revenue.assumptions_summary:
        _add_bullets(document, revenue.assumptions_summary[:8])
    if revenue.invariant_violations:
        _add_heading(document, "Model Warnings", 2)
        _add_bullets(document, revenue.invariant_violations)

    _add_heading(document, "Execution Roadmap", 1)
    _add_table(
        document,
        ["Phase", "Duration", "Focus", "Key milestones / success metrics"],
        [
            [
                f"{phase.phase_number}. {phase.name}",
                f"{phase.duration_weeks} weeks",
                phase.focus,
                "\n".join([f"- {item}" for item in phase.key_milestones] + ["Success metrics:"] + [f"- {item}" for item in phase.success_metrics]),
            ]
            for phase in roadmap.execution_phases
        ],
    )

    _add_heading(document, "Top Cross-Functional Risks", 1)
    _add_table(
        document,
        ["Category", "Risk", "Likelihood", "Impact", "Mitigation"],
        [
            [risk.category, risk.risk, risk.likelihood, risk.impact, risk.mitigation]
            for risk in roadmap.top_risks
        ],
    )

    _add_heading(document, "90-Day GTM Milestone Calendar", 1)
    document.add_paragraph(marketing.positioning_statement)
    _add_table(
        document,
        ["Week", "Focus", "Target KPI", "Deliverables"],
        [
            [
                f"W{milestone.week}" if milestone.week else f"M{milestone.month}",
                milestone.focus,
                milestone.target_kpi,
                "; ".join(milestone.deliverables) or "-",
            ]
            for milestone in marketing.gtm_90_day_plan
        ],
    )
    if marketing.prioritized_next_actions:
        _add_heading(document, "Prioritized Next Actions", 1)
        _add_bullets(document, marketing.prioritized_next_actions[:8])

    document.add_paragraph()
    disclaimer = document.add_paragraph(roadmap.disclaimer)
    disclaimer.runs[0].italic = True
    disclaimer.runs[0].font.size = Pt(8)
    disclaimer.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_roadmap_docx(roadmap: StartupRoadmap) -> bytes:
    """Backward-friendly alias for the Word exporter."""
    return generate_docx(roadmap)


generate_docx_report = generate_docx
