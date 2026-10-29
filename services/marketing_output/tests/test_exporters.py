"""Tests for roadmap synthesis and PDF / DOCX / HTML export engines."""

from io import BytesIO
import zipfile

from docx import Document

from services.marketing_output.app.export.docx_exporter import generate_docx
from services.marketing_output.app.export.pdf_exporter import generate_pdf
from services.marketing_output.app.export.roadmap_renderer import render_roadmap_html
from services.marketing_output.app.interface import run_roadmap_synthesis
from services.marketing_output.tests.factories import build_sample_state


def test_roadmap_synthesis_populates_phases_and_risks():
    roadmap = run_roadmap_synthesis(build_sample_state())

    assert [phase.name for phase in roadmap.execution_phases] == ["Validation", "MVP", "Launch", "Scale"]
    assert all(phase.duration_weeks > 0 for phase in roadmap.execution_phases)
    assert len(roadmap.top_risks) == 5
    assert all(risk.mitigation for risk in roadmap.top_risks)
    assert 0.0 <= roadmap.overall_confidence_score <= 1.0
    assert "Overall confidence" in roadmap.executive_summary
    assert "SOC2" in " ".join(risk.risk for risk in roadmap.top_risks)


def test_pdf_exporter_produces_valid_pdf_bytes():
    content = generate_pdf(run_roadmap_synthesis(build_sample_state()))

    assert content.startswith(b"%PDF")
    assert b"%%EOF" in content[-1024:]
    assert len(content) > 5000


def test_docx_exporter_produces_valid_editable_document():
    content = generate_docx(run_roadmap_synthesis(build_sample_state()))

    assert content.startswith(b"PK")
    assert zipfile.is_zipfile(BytesIO(content))
    document = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Enterprise CRM workflow automation platform" in text
    assert "Business Model Canvas" in text
    assert len(document.tables) >= 5


def test_html_renderer_produces_standalone_responsive_report():
    content = render_roadmap_html(run_roadmap_synthesis(build_sample_state()))

    assert content.startswith("<!DOCTYPE html>")
    assert '<meta name="viewport"' in content
    assert "Business Model Canvas" in content
    assert "Top Cross-Functional Risks" in content
    assert "90-Day GTM Milestone Calendar" in content
    assert "Enterprise CRM workflow automation platform" in content
